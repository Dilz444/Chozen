from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "products-template.xlsx")

IVORY = "FBF8F4"
IVORY2 = "F4ECE4"
BLUSH = "EED9D6"
INK = "2A2225"
INK2 = "62565A"
ROSE = "94495D"
LINE = "E6DCD3"
GREY = "8A8185"
FONT = "Calibri"

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_BORDER = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=ROSE))
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
HEAD_FILL = PatternFill("solid", fgColor=IVORY)
EX_FILL = PatternFill("solid", fgColor="F7F4F1")
HEAD_FONT = Font(name=FONT, bold=True, color=INK, size=11)
BODY_FONT = Font(name=FONT, color=INK, size=11)
EX_FONT = Font(name=FONT, italic=True, color=GREY, size=11)


def note(text, w=320, h=180):
    c = Comment(text, "ChoZen website builder")
    c.width, c.height = w, h
    return c


wb = Workbook()

# ---------------------------------------------------------------- Lists
lists = wb.create_sheet("Lists")
LISTS = {
    "Line": ["Fresh flowers", "Balloons", "Flowers + balloons", "Faux flowers", "Crystal",
             "Jewellery", "Other gift", "Subscription", "Hire"],
    "Stone treatment": ["Unknown", "None known", "Dyed", "Heated", "Stabilised",
                        "Reconstituted", "Synthetic", "Other"],
    "Delivery type": ["Local same-day only", "UK shipping", "Both"],
    "Yes / No": ["Yes", "No"],
    "Occasion ideas (type these in, comma-separated)": [
        "Birthday", "Anniversary", "New baby", "Thank you", "Get well", "Sympathy",
        "Mother's Day", "Valentine's", "Just because"],
    "Sole trader or Ltd": ["Sole trader", "Limited company", "Not sure"],
}
ranges = {}
for ci, (title, vals) in enumerate(LISTS.items(), start=1):
    col = get_column_letter(ci)
    h = lists.cell(row=1, column=ci, value=title)
    h.font, h.fill, h.border, h.alignment = HEAD_FONT, HEAD_FILL, HEAD_BORDER, WRAP_TOP
    for ri, v in enumerate(vals, start=2):
        c = lists.cell(row=ri, column=ci, value=v)
        c.font, c.border = BODY_FONT, BORDER
    ranges[title] = f"Lists!${col}$2:${col}${len(vals) + 1}"
    lists.column_dimensions[col].width = 28
lists.row_dimensions[1].height = 34
lists.freeze_panes = "A2"
note_cell = lists.cell(row=13, column=1,
                       value="These lists feed the dropdowns on the Products sheet. Please don't edit them. "
                             "If something's missing, add it under \"Notes for the builder\" instead.")
note_cell.font = Font(name=FONT, italic=True, color=INK2)
lists.merge_cells("A13:D14")
note_cell.alignment = WRAP_TOP

# ---------------------------------------------------------------- Products
ws = wb.create_sheet("Products")
# (header, width, note, list-name or None)
COLS = [
    ("Product name", 26,
     "Your name for it, the way you'd say it in the shop.\n\ne.g. \"Blush rose bouquet\" or \"Amethyst cluster\".\n\nOne row per product. If it comes in sizes, keep it on one row and put the sizes in the \"Sizes / options\" column.", None),
    ("Line", 18,
     "Pick one from the dropdown:\nFresh flowers, Balloons, Flowers + balloons, Faux flowers, Crystal, Jewellery, Other gift, Subscription, Hire.", "Line"),
    ("Category", 20,
     "Anything that groups it, in your own words.\n\ne.g. hand-tied bouquet, vase arrangement, hatbox, balloon bundle, crystal cluster, tumble stone, bracelet, candle.", None),
    ("Occasions", 24,
     "Which occasions it suits. Type them with commas between.\n\nIdeas: Birthday, Anniversary, New baby, Thank you, Get well, Sympathy, Mother's Day, Valentine's, Just because.\n\nTip: please don't put balloons under Sympathy.", None),
    ("Price £", 11,
     "The starting price in pounds, numbers only (e.g. 45).\n\nIf it has sizes, put the smallest size's price here and the rest in the next column.", None),
    ("Sizes / options and their prices", 26,
     "Each size or option with its price.\n\ne.g. \"Standard £45 / Large £65\" or \"Bracelet: small £20 / medium £20 / large £22\".\n\nLeave blank if there's only one.", None),
    ("Short description in your words", 42,
     "1 to 3 sentences, how you'd describe it to a customer in the shop. Don't worry about polishing it, we'll write the final version from this.\n\nFor crystals: please don't say it heals or helps with anything. Say what it's traditionally associated with instead.", None),
    ("What's in it", 36,
     "Everything that arrives.\n\nFlowers: each flower type, how many stems, and colours (e.g. 10 blush roses, 5 white spray roses, eucalyptus).\n\nBalloons: how many, latex or foil, colours, and whether they arrive filled with helium (yes/no).\n\nGifts: every item in the box.", None),
    ("Size", 22,
     "Roughly how big, in cm: height × width.\n\nCrystals: weight in grams plus size (e.g. about 6 × 5 cm, 150 g).\n\nJewellery: length or ring size.", None),
    ("Materials (metal, stone)", 28,
     "Mainly for jewellery and gifts.\n\nMetal and its standard, exactly as your supplier states it: e.g. 925 sterling silver, gold-plated brass, stainless steel.\n\nThe stone: e.g. rose quartz, amethyst, howlite.\n\nIf you're not sure, write \"not sure\". We'd rather check than guess.", None),
    ("Stone treatment", 17,
     "Has the stone been treated? Pick from the dropdown: Unknown, None known, Dyed, Heated, Stabilised, Reconstituted, Synthetic, Other.\n\nYour supplier's listing or invoice usually says. \"Unknown\" is fine, we'll ask them.\n\nLeave blank for flowers and balloons.", "Stone treatment"),
    ("Who it's for", 22,
     "Who usually buys it, or who it's for.\n\ne.g. mums, new parents, a friend who's poorly, a partner, the office.", None),
    ("Stock", 14,
     "How many you have now (a number), or type \"made to order\".", None),
    ("Delivery type", 20,
     "Pick one:\n• Local same-day only: you hand-deliver it (fresh flowers, helium balloons).\n• UK shipping: it can go in the post.\n• Both: posted UK-wide, and you can also take it out with the local deliveries.", "Delivery type"),
    ("Weight (g, packed)", 14,
     "Weight in grams once it's boxed and ready to post. Kitchen scales are fine.\n\nOnly needed for things that go in the post.", None),
    ("Contains lilies?", 12,
     "Yes or No. Lilies are dangerous to cats, so the website shows a warning on these.", "Yes / No"),
    ("Latex balloons?", 12,
     "Yes or No. Latex balloons need a choking warning for under-8s and a \"natural rubber latex\" note on the page.", "Yes / No"),
    ("Personalised / made to their spec?", 16,
     "Yes if it's made specially for that customer (a name or age on a balloon, chosen colours, an engraved item). It changes the returns rules, so please be honest here.", "Yes / No"),
    ("Photo filenames", 32,
     "The names of the photos for this product: product-name-1.jpg (the main square photo), then -2, -3.\n\ne.g. blush-rose-bouquet-1.jpg, blush-rose-bouquet-2.jpg, blush-rose-bouquet-3.jpg\n\nSee the photo guide.", None),
    ("Notes for the builder", 34,
     "Anything else: seasonal changes, what you'd swap if a flower's not available, a supplier name, questions for us.", None),
]
HEADERS = [c[0] for c in COLS]
NCOL = len(COLS)

for ci, (head, width, txt, _) in enumerate(COLS, start=1):
    c = ws.cell(row=1, column=ci, value=head)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, HEAD_BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.comment = note(txt)
    ws.column_dimensions[get_column_letter(ci)].width = width
ws.row_dimensions[1].height = 48
ws.freeze_panes = "B2"

EXAMPLES = [
    {
        "Product name": "EXAMPLE: Blush rose bouquet",
        "Line": "Fresh flowers",
        "Category": "Hand-tied bouquet",
        "Occasions": "Birthday, Anniversary, Thank you, Just because",
        "Price £": "£—",
        "Sizes / options and their prices": "Standard £— / Large £—",
        "Short description in your words": "Soft pink roses with eucalyptus, hand-tied in ivory paper with a cream ribbon. My go-to when someone says \"something pretty, not too loud\".",
        "What's in it": "Standard: 10 blush roses, 5 white spray roses, eucalyptus, white waxflower. Large: 15 blush roses, 7 white spray roses, more eucalyptus.",
        "Size": "Standard about 45 × 30 cm. Large about 55 × 40 cm.",
        "Materials (metal, stone)": "",
        "Stone treatment": "",
        "Who it's for": "Mums, partners, friends",
        "Stock": "made to order",
        "Delivery type": "Local same-day only",
        "Weight (g, packed)": "",
        "Contains lilies?": "No",
        "Latex balloons?": "No",
        "Personalised / made to their spec?": "No",
        "Photo filenames": "blush-rose-bouquet-1.jpg, blush-rose-bouquet-2.jpg, blush-rose-bouquet-3.jpg",
        "Notes for the builder": "If blush roses aren't in, I use pale peach. Comes in a water bubble.",
    },
    {
        "Product name": "EXAMPLE: Birthday balloon bundle",
        "Line": "Balloons",
        "Category": "Balloon bundle",
        "Occasions": "Birthday, New baby",
        "Price £": "£—",
        "Sizes / options and their prices": "5 balloons £— / 8 balloons £—",
        "Short description in your words": "A bunch of pink, white and gold balloons with a foil number for their age, tied to a weight so it stands up on its own.",
        "What's in it": "4 latex balloons (pink, white, gold confetti) + 1 foil number balloon (gold). Helium-filled: yes. Comes with a weight.",
        "Size": "About 150 cm tall once filled",
        "Materials (metal, stone)": "",
        "Stone treatment": "",
        "Who it's for": "Kids' and grown-ups' birthdays",
        "Stock": "made to order",
        "Delivery type": "Local same-day only",
        "Weight (g, packed)": "",
        "Contains lilies?": "No",
        "Latex balloons?": "Yes",
        "Personalised / made to their spec?": "Yes",
        "Photo filenames": "birthday-balloon-bundle-1.jpg, birthday-balloon-bundle-2.jpg",
        "Notes for the builder": "Customer picks the number (0–9) and the colours. Roughly how long they stay up: [your experience].",
    },
    {
        "Product name": "EXAMPLE: Amethyst cluster",
        "Line": "Crystal",
        "Category": "Crystal cluster",
        "Occasions": "Birthday, Thank you, Just because",
        "Price £": "£—",
        "Sizes / options and their prices": "Small £— / Medium £—",
        "Short description in your words": "A deep purple amethyst cluster with sparkly points. Amethyst is traditionally linked with calm, and people love it on a desk or bedside table.",
        "What's in it": "1 amethyst cluster in a gift box, with a little card about the stone",
        "Size": "Small about 6 × 5 cm, around 150 g. Medium about 9 × 7 cm, around 300 g.",
        "Materials (metal, stone)": "Natural amethyst. Origin: Brazil (from the supplier's invoice)",
        "Stone treatment": "None known",
        "Who it's for": "Friends, teachers, anyone who likes crystals",
        "Stock": "4 small, 2 medium",
        "Delivery type": "Both",
        "Weight (g, packed)": "Small 250 / Medium 450",
        "Contains lilies?": "No",
        "Latex balloons?": "No",
        "Personalised / made to their spec?": "No",
        "Photo filenames": "amethyst-cluster-1.jpg, amethyst-cluster-2.jpg, amethyst-cluster-3.jpg",
        "Notes for the builder": "Every one is slightly different. The customer gets one like the photo.",
    },
]

FIRST_EMPTY = 2 + len(EXAMPLES)
LAST_ROW = FIRST_EMPTY + 60 - 1  # 60 empty rows

for ri, ex in enumerate(EXAMPLES, start=2):
    for ci, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=ri, column=ci, value=ex[h] if ex[h] != "" else None)
        c.font, c.fill, c.border, c.alignment = EX_FONT, EX_FILL, BORDER, WRAP_TOP
    ws.row_dimensions[ri].height = 96

price_col = HEADERS.index("Price £") + 1
for ri in range(FIRST_EMPTY, LAST_ROW + 1):
    for ci in range(1, NCOL + 1):
        c = ws.cell(row=ri, column=ci)
        c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP_TOP
        if ci == price_col:
            c.number_format = '"£"#,##0.00'
    ws.row_dimensions[ri].height = 42

# Dropdowns (applied a little beyond the formatted rows so extra rows still work)
VAL_END = 500
for ci, (head, _, _, listname) in enumerate(COLS, start=1):
    if not listname:
        continue
    col = get_column_letter(ci)
    dv = DataValidation(type="list", formula1=f"{ranges[listname]}", allow_blank=True,
                        showDropDown=False)  # False = show the arrow (openpyxl naming is inverted)
    dv.error = "Please pick one from the list. If none fits, choose the nearest and explain in Notes for the builder."
    dv.errorTitle = "Pick from the list"
    dv.prompt = "Tap the arrow and pick one."
    dv.promptTitle = head
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{VAL_END}")

# Gentle number check on price (a number, or blank)
pcol = get_column_letter(price_col)
dv_price = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv_price.errorTitle = "Price"
dv_price.error = "Please type the price as a number only, e.g. 45. Put sizes and their prices in the next column."
dv_price.showErrorMessage = True
ws.add_data_validation(dv_price)
dv_price.add(f"{pcol}{FIRST_EMPTY}:{pcol}{VAL_END}")

ws.auto_filter.ref = f"A1:{get_column_letter(NCOL)}{LAST_ROW}"
ws.sheet_view.zoomScale = 100

# ---------------------------------------------------------------- Delivery & business facts
fx = wb.create_sheet("Delivery & business facts")
FX_COLS = [("Topic", 20), ("Question", 46), ("Why we ask", 40), ("Your answer", 48), ("Ref", 6)]
for ci, (h, w) in enumerate(FX_COLS, start=1):
    c = fx.cell(row=1, column=ci, value=h)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, HEAD_BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    fx.column_dimensions[get_column_letter(ci)].width = w
fx["D1"].comment = note("Type your answer in this column. \"Not sure\" or \"not yet\" is a perfectly good answer.", 260, 80)
fx.row_dimensions[1].height = 30
fx.freeze_panes = "A2"

FACTS = [
    ("Your business", None, None, None),
    ("Address", "Where are you based? Can customers visit a shop or studio, or do you work from home? Full address with postcode.", "The law says the website must show a real address (not a PO box). If you work from home we can keep it off Google Maps.", "B1"),
    ("Legal name", "Your full legal name, or the company's registered name if it's a limited company.", "It goes in the small print at the bottom of every page.", "B2"),
    ("Sole trader or Ltd", "Are you a sole trader or a limited company? If limited, the company number.", "Limited companies have to show their number on the website.", "B2"),
    ("VAT", "Are you VAT registered? If yes, your VAT number.", "Changes how prices and invoices are shown.", "B2"),
    ("Opening hours", "Your opening hours for each day (or the hours you answer the phone / take orders).", "Shown on Google and the contact page.", "B14"),
    ("Local same-day delivery", None, None, None),
    ("Delivery postcodes", "Which postcodes or areas do you deliver to yourself? e.g. EN1, EN2, EN3, N9, N13, N14, N18, N21. Anywhere you'd go further for a bigger order?", "We'll only promise the areas you list.", "B3"),
    ("Fee per area", "The delivery charge for each area. Is there a minimum order for delivery?", "The fee has to show before people get to the checkout.", "B4"),
    ("Delivery days", "Which days do you deliver? Sundays? Bank holidays?", "So the website never promises a day you're closed.", "B5"),
    ("Cut-off time", "What's the latest time someone can order for same-day delivery? The brief says 10am. Could you do 12 noon or 1pm?", "The biggest thing that wins local orders. Other florists nearby take orders until 2pm or later.", "B6"),
    ("Delivery window", "Roughly what hours do deliveries arrive (e.g. 10am to 6pm)? Can customers choose morning or afternoon?", "People want to know when to be home.", "B7"),
    ("Balloons", "Which balloons do you use (latex, foil, bubble)? Are they filled with helium when delivered? Roughly how long do they stay up?", "For the product pages and the safety notes.", "B8"),
    ("UK shipping (posted gifts)", None, None, None),
    ("Shipping service", "Which service do you use to post things (Royal Mail, Evri, other)? Tracked?", "Shown on the delivery page.", "—"),
    ("Shipping price", "What do you charge for UK postage? Is there a spend where postage is free?", "Shown before checkout.", "—"),
    ("Dispatch time", "How soon do you post after an order (same day, next working day)?", "Shown on the product and delivery pages.", "—"),
    ("Subscription", None, None, None),
    ("What it is", "What do subscribers get? Fresh flowers? Weekly, fortnightly or monthly? Local delivery only?", "For the subscription page.", "B9"),
    ("Price and terms", "The price, and how someone pauses or cancels.", "Subscription rules must be clear by law.", "B9"),
    ("Floral hire", None, None, None),
    ("What you hire out", "Faux or fresh flowers? For offices, restaurants, events, weddings?", "For the hire page.", "B10"),
    ("Price and contract", "Monthly price (or 'from' price), how long the contract is, and how often you refresh or swap the flowers.", "So businesses know what to expect before they enquire.", "B10"),
    ("Buy-back programme", None, None, None),
    ("What and how", "What do you buy back (crystals? faux arrangements?), how much do you pay, and what condition must it be in?", "It's a promise to customers, so it needs clear terms.", "B11"),
    ("Anything else", None, None, None),
    ("Crystal sourcing", "Who are your crystal suppliers and which countries do the stones come from? Any paperwork that shows ethical sourcing?", "We only say what you can show.", "B12"),
    ("Reviews", "Do you have reviews anywhere (Google, Facebook, Etsy)? Links please.", "We only show real reviews.", "B16"),
    ("\"Claims Payment\"", "What is the \"Claims Payment\" link on your current website for?", "So we put it in the right place.", "B18"),
    ("Current orders", "Does your current website take real orders? Any customer list or newsletter sign-ups?", "So nothing gets lost when we switch.", "B20"),
]
r = 2
answer_rows = {}
for topic, q, why, ref in FACTS:
    if q is None:
        c = fx.cell(row=r, column=1, value=topic)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        for ci in range(1, 6):
            fx.cell(row=r, column=ci).fill = PatternFill("solid", fgColor=ROSE)
        fx.row_dimensions[r].height = 22
        fx.cell(row=r, column=1).alignment = Alignment(vertical="center")
    else:
        vals = [topic, q, why, None, ref]
        for ci, v in enumerate(vals, start=1):
            c = fx.cell(row=r, column=ci, value=v)
            c.border, c.alignment = BORDER, WRAP_TOP
            c.font = Font(name=FONT, color=INK, bold=(ci == 1))
            if ci == 3:
                c.font = Font(name=FONT, color=INK2, italic=True)
            if ci == 5:
                c.font = Font(name=FONT, color=GREY, size=9)
        fx.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FFFFFF")
        fx.row_dimensions[r].height = 64
        answer_rows[topic] = r
    r += 1

dv_st = DataValidation(type="list", formula1=f"{ranges['Sole trader or Ltd']}", allow_blank=True)
dv_st.prompt = "Pick one, then add the company number after it if limited (you can type over the choice)."
dv_st.showErrorMessage = False
fx.add_data_validation(dv_st)
dv_st.add(f"D{answer_rows['Sole trader or Ltd']}")

# ---------------------------------------------------------------- How to fill this in
g = wb.active
g.title = "How to fill this in"
g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 3
g.column_dimensions["B"].width = 30
g.column_dimensions["C"].width = 70
g.column_dimensions["D"].width = 3

row = 2


def put(text, font=None, height=None, merge=True, fill=None):
    global row
    c = g.cell(row=row, column=2, value=text)
    c.font = font or BODY_FONT
    c.alignment = WRAP_TOP
    if merge:
        g.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    if fill:
        for ci in (2, 3):
            g.cell(row=row, column=ci).fill = fill
    if height:
        g.row_dimensions[row].height = height
    row += 1


TITLE = Font(name="Georgia", size=20, color=ROSE)
H2 = Font(name="Georgia", size=14, color=ROSE, bold=False)
SMALL = Font(name=FONT, size=10, italic=True, color=INK2)

put("Your products, in your words", TITLE, 32)
put("Hello Beyzan. This sheet is how your products get onto the new website. You tell us what each one is, "
    "in your own words, and we write the final descriptions from your notes. No need to make it sound polished. "
    "Short notes are perfect.", height=50)
row += 1
put("How to fill it in", H2, 22)
steps = [
    "1.  Open the Products tab. Each row is one product.",
    "2.  The first three rows are grey EXAMPLES. Have a look, then start on the first white row underneath. You can leave the examples where they are.",
    "3.  Tap any heading to see a short note on what goes there (on a phone, tap the heading cell; on a laptop, hover over the little red corner).",
    "4.  Some boxes have a dropdown. Tap the cell, then the little arrow, and pick one.",
    "5.  Not sure about something? Leave it blank or write \"not sure\". We'll ask you. Please don't guess metals, stone treatments or prices.",
    "6.  If a product comes in sizes, keep it on one row and list the sizes and prices in \"Sizes / options\".",
    "7.  Then fill in the \"Delivery & business facts\" tab: your delivery areas, fees, times and so on.",
    "8.  Take the photos (see the separate photo guide) and name them to match the \"Photo filenames\" column.",
]
for s in steps:
    put(s, height=34)
row += 1
put("One example, column by column", H2, 22)
put("Here's the first example row (the blush rose bouquet) with what each answer tells us.", height=20)

hdr_fill = PatternFill("solid", fgColor=IVORY)
for ci, v in ((2, "Column and example answer"), (3, "Why this helps")):
    c = g.cell(row=row, column=ci, value=v)
    c.font, c.fill, c.border, c.alignment = HEAD_FONT, hdr_fill, HEAD_BORDER, WRAP_TOP
row += 1
ex = EXAMPLES[0]
why = {
    "Product name": "Your name for it. We'll use it to make the web address and the page title.",
    "Line": "Puts it in the right part of the shop (here: same-day flowers).",
    "Category": "Helps us group similar things together.",
    "Occasions": "It'll show up in the Birthday, Anniversary and Thank you sections too.",
    "Price £": "The real examples show £— because we never make up prices. You type the real number, e.g. 45.",
    "Sizes / options and their prices": "Becomes the size choice on the product page.",
    "Short description in your words": "We turn this into the final description, keeping your way of saying things.",
    "What's in it": "The most useful bit. People want to know exactly what arrives, so nobody is surprised at the door.",
    "Size": "Shown as \"about 45cm tall\" so the customer can picture it.",
    "Who it's for": "Helps us write gift ideas and link it from the right pages.",
    "Stock": "Made to order means it never shows as sold out.",
    "Delivery type": "Fresh flowers are hand-delivered locally only, so the website won't offer to post them.",
    "Contains lilies?": "No here. If Yes, the page warns cat owners.",
    "Latex balloons?": "No here. If Yes, the page adds the balloon safety note.",
    "Personalised / made to their spec?": "No here, so the normal returns rules apply.",
    "Photo filenames": "Tells us which photos belong to this product. The -1 photo is the main square one.",
    "Notes for the builder": "Anything we should know. Here: what you'd swap in if blush roses aren't available.",
}
for h in HEADERS:
    if h not in why:
        continue
    val = ex[h] or "(blank)"
    b = g.cell(row=row, column=2, value=f"{h}\n{val}")
    b.font, b.alignment, b.border = BODY_FONT, WRAP_TOP, BORDER
    c = g.cell(row=row, column=3, value=why[h])
    c.font, c.alignment, c.border = BODY_FONT, WRAP_TOP, BORDER
    lines = max(2 + len(str(val)) // 30, 1 + len(why[h]) // 68)
    g.row_dimensions[row].height = max(34, 15 * lines)
    row += 1
put("Blank columns (Materials, Stone treatment, Weight) are for jewellery, crystals and things that get posted. "
    "Fine to leave empty for fresh flowers.", SMALL, 30)
row += 1
put("For crystals and jewellery", H2, 22)
for s in [
    "•  Please don't write that a crystal heals or helps with anything (sleep, anxiety, energy). We're not allowed to say it. \"Traditionally associated with\" is fine.",
    "•  Only call jewellery silver or gold if your supplier says so. Plated items are \"gold-plated\" or \"silver-plated\".",
    "•  If a stone's been dyed or treated, or you don't know, say so in Stone treatment. \"Unknown\" is fine.",
]:
    put(s, height=34)
row += 1
put("What happens next", H2, 22)
for s in [
    "1.  Send the spreadsheet back with your photos (by email or a shared folder; WhatsApp shrinks photos).",
    "2.  We write each product's name, description and details from your notes.",
    "3.  You check them on a preview of the new shop and tell us what to change.",
    "4.  Once you're happy, they go live.",
]:
    put(s, height=22)
row += 1
put("Thank you. Questions at any point, just ask.", Font(name="Georgia", size=12, italic=True, color=ROSE), 20)

for sheet in (g, ws, fx, lists):
    sheet.sheet_properties.tabColor = ROSE
lists.sheet_properties.tabColor = LINE
wb.move_sheet("Lists", offset=10)  # Lists last
wb.active = 0
wb.save(OUT)
print("saved", OUT, "rows", FIRST_EMPTY, LAST_ROW)
